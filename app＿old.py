from flask import Flask, render_template, request, send_file
import openpyxl, os, json
from datetime import datetime
from openpyxl.styles import PatternFill

app = Flask(__name__)

# 讀取設定檔
with open('config.json', 'r', encoding='utf-8') as f:
    config = json.load(f)

TEMPLATE_PATH = config["input_file"]
OUTPUT_FOLDER = config["output_folder"]

REASON_OPTIONS = [
    "體檢", "休假", "年休返泰", "事假返鄉",
    "工傷", "曠職", "病假", "待返",
    "遣返", "提前解聘", "逃跑", "調派"
]

def update_excel(absentees):
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws_main = wb["出勤表"]
    ws_log = wb["休假調查表(新)"]

    reason_map = {emp_id.strip(): reason for emp_id, reason in absentees}
    emp_columns = [2, 8, 14, 20, 26]
    start_row, end_row = 6, 61

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    red_fill = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")

    count_map = {key: 0 for key in REASON_OPTIONS}

    for row in range(start_row, end_row + 1):
        for col in emp_columns:
            emp_cell = ws_main.cell(row=row, column=col)
            upper_cell = ws_main.cell(row=row, column=col + 1)
            emp_id = str(emp_cell.value).strip() if emp_cell.value else ""
            if len(emp_id) < 3:
                continue

            if emp_id in reason_map:
                reason = reason_map[emp_id]
                upper_cell.value = "X"
                emp_cell.fill = red_fill if "曠職" in reason or "逃跑" in reason else yellow_fill
                if reason in count_map:
                    count_map[reason] += 1
            else:
                upper_cell.value = "V"

    # 寫入休假調查表
    insert_row = 5
    serial_number = 1  # 從 1 開始的序號
    today = datetime.now().strftime("%m/%d")

    for emp_id, reason in absentees:
        ws_log.cell(row=insert_row, column=1).value = today           # 第一欄：當天日期
        ws_log.cell(row=insert_row, column=2).value = serial_number   # 第二欄：序號
        ws_log.cell(row=insert_row, column=3).value = "GC01"          # 第三欄：固定填 GC01
        ws_log.cell(row=insert_row, column=4).value = emp_id          # 第四欄：工號
        ws_log.cell(row=insert_row, column=5).value = reason          # 第五欄：未出工原因
        ws_log.cell(row=insert_row, column=6).value = "宿舍"          # 第六欄：固定填 宿舍

        insert_row += 1
        serial_number += 1

    # 寫入統計欄位
    ws_main["C62"].value = count_map["體檢"]
    ws_main["C63"].value = count_map["工傷"]
    ws_main["C64"].value = count_map["遣返"]
    ws_main["K62"].value = count_map["休假"]
    ws_main["K63"].value = count_map["曠職"]
    ws_main["K64"].value = count_map["提前解聘"]
    ws_main["T62"].value = count_map["年休返泰"]
    ws_main["T63"].value = count_map["病假"]
    ws_main["T64"].value = count_map["逃跑"]
    ws_main["AA62"].value = count_map["事假返鄉"]
    ws_main["AA63"].value = count_map["待返"]
    ws_main["AA64"].value = count_map["調派"]

    today_str = datetime.today().strftime("%Y-%m-%d")
    output_file = f"{OUTPUT_FOLDER}/每天出工統計表_{today_str}.xlsx"
    wb.save(output_file)
    return output_file

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        emp_ids = request.form.getlist('emp_id')
        reasons = request.form.getlist('reason')
        absentees = [(emp_ids[i], reasons[i]) for i in range(len(emp_ids)) if emp_ids[i].strip()]
        result_path = update_excel(absentees)
        return send_file(result_path, as_attachment=True)
    return render_template('index.html', reasons=REASON_OPTIONS)

if __name__ == '__main__':
    app.run(debug=True)
